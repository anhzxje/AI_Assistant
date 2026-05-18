"""CSV/JSON/Excel tabular parser used by file ingestion jobs."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any
from zipfile import ZipFile
import xml.etree.ElementTree as ET


SUPPORTED_EXTENSIONS = {".csv", ".json", ".jsonl", ".xlsx"}


def parse_tabular_file(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return parse_csv(path)
    if suffix == ".json":
        value = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(value, list):
            return [row for row in value if isinstance(row, dict)]
        if isinstance(value, dict):
            return [value]
        raise ValueError("JSON tabular input must be an object or list of objects.")
    if suffix == ".jsonl":
        return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    if suffix == ".xlsx":
        return parse_xlsx(path)
    raise ValueError(f"Unsupported tabular extension: {suffix}")


def parse_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def parse_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:  # noqa: BLE001 - dependency is optional.
        return parse_xlsx_without_openpyxl(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        records.extend(rows_to_records(rows, sheet_name=sheet.title))
    return records


def rows_to_records(rows: list[tuple[Any, ...]] | list[list[Any]], *, sheet_name: str | None = None) -> list[dict[str, Any]]:
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    records: list[dict[str, Any]] = []
    for row in rows[1:]:
        record = {headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]}
        if sheet_name:
            record["sheet_name"] = sheet_name
        records.append(record)
    return records


def parse_xlsx_without_openpyxl(path: Path) -> list[dict[str, Any]]:
    with ZipFile(path) as archive:
        sheet_names = workbook_sheet_names(archive)
        sheet_paths = sorted(name for name in archive.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"))
        records: list[dict[str, Any]] = []
        for index, sheet_path in enumerate(sheet_paths):
            sheet_name = sheet_names[index] if index < len(sheet_names) else Path(sheet_path).stem
            rows = worksheet_rows(archive.read(sheet_path))
            records.extend(rows_to_records(rows, sheet_name=sheet_name))
        return records


def workbook_sheet_names(archive: ZipFile) -> list[str]:
    try:
        root = ET.fromstring(archive.read("xl/workbook.xml"))
    except Exception:
        return []
    ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    return [str(sheet.attrib.get("name") or "") for sheet in root.findall(".//main:sheet", ns)]


def worksheet_rows(raw_xml: bytes) -> list[list[str]]:
    root = ET.fromstring(raw_xml)
    ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rows: list[list[str]] = []
    for row in root.findall(".//main:row", ns):
        values: list[str] = []
        for cell in row.findall("main:c", ns):
            text_node = cell.find(".//main:t", ns)
            values.append(text_node.text if text_node is not None and text_node.text is not None else "")
        rows.append(values)
    return rows
